import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

FONT="Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT, size=10)
RED = PatternFill("solid", fgColor="F8CBAD")
GREEN = PatternFill("solid", fgColor="C6E0B4")
YELLOW = PatternFill("solid", fgColor="FFF2CC")

wb = load_workbook("/sessions/gallant-sharp-hopper/mnt/GM Improvement/GM_V4_Model_Aug2026.xlsx")
ws = wb.create_sheet("VIP Gold ASP Strategy", index=3)  # place right after Price-Coupon Stacking Risk
ws.sheet_view.showGridLines = False
ws.column_dimensions['A'].width = 100

def h1(cell, text, size=13, color="1F4E78"):
    cell.value = text; cell.font = Font(name=FONT, size=size, bold=True, color=color)
def body(cell, text, italic=False):
    cell.value = text; cell.font = Font(name=FONT, size=11, italic=italic)
    cell.alignment = Alignment(wrap_text=True, vertical='top')

r=1
h1(ws.cell(row=r,column=1), "VIP Gold ASP Strategy -- sourced from external deepdive (Aug MTD, Jun-Aug WoW), fact-checked and connected to V4"); r+=1
body(ws.cell(row=r,column=1), "Origin: a separate 'VIP Gold users deepdive' analysis (11-week WoW, Jun 1 - Aug 10) was fact-checked against this model's raw booking data (Jul26-Aug2 window) and cross-referenced with the Price, Coupon, and Segment tabs below.", italic=True); r+=2

h1(ws.cell(row=r,column=1), "Fact-check verdict", size=12); r+=1
for s in [
 "A prior memo ('Gemini memo') claimed August VIP transactions are operationally unprofitable, based on treating diagnostic_cost/express_slot/report_hardcopy_cost as company costs. They are customer-paid, waivable FEES, not costs -- cpt is the only real cost field (established in V4 Methodology section 2). That memo's 'Phlebo Revolt' framing is REJECTED.",
 "Independently verified on this model's own Jul26-Aug2 raw data: VIP-discount-only (isolated, no coupon stacked) bookings run 70.96% GM% -- healthy, not unprofitable. The deepdive's own number (76.6%, broader Aug MTD window) points the same direction; the ~6pp gap is a window/definition difference, not a disagreement.",
 "Also independently confirmed: RedCash burn in VIP Gold is ~Rs 0 (exact match), members/booking trending up (1.213 in this window vs the deepdive's 1.17->1.22 WoW trend), and CPT% of price for VIP+coupon stacked bookings at 24.9% (deepdive: 24.5-26.7% range).",
]:
    body(ws.cell(row=r,column=1), s); r+=1
r+=1

h1(ws.cell(row=r,column=1), "What V4 adds that the deepdive couldn't see", size=12); r+=1
for s in [
 "FINDING A -- Ancillary attach is channel-dependent and Phlebo is a dead zone: for VIP Gold specifically, Phlebo-channel bookings run 0% express-slot AND 0% hardcopy attach (vs Self 4.1%/1.0%, Sales 3.4%/0.4%). The deepdive's #1 recommendation (ancillary fee push) should start at the Self/app cart stage where some attach already exists to build on -- Phlebo needs a different, separate fix (or a check on whether the offer is even technically available at that touchpoint).",
 "FINDING B -- 5 of the 6 Price-lever 'stacking risk' packages are VIP-Gold-heavy: PL124I (48% VIP Gold share), PL08I (35%), PL334 (33%), PL91 (37%), PL203 (34%) all appear on the Price-Coupon Stacking Risk tab (already-discounted packages flagged for a price cut). A price cut on these doesn't just risk generic coupon stacking -- it specifically compounds against the segment this strategy is trying to protect margin on. See Action #6 below.",
 "FINDING C -- V4's own Segment Diagnostic already showed VIP Gold gets meaningfully SHALLOWER coupons than Regular (8.7-10.1% vs ~20% depth) -- structurally consistent with the deepdive's caution against a hard discount cliff: VIP Gold's margin problem was never primarily discount depth, which supports prioritizing the low-risk ancillary/basket levers (deepdive's #1, #2) well ahead of the discount-tiering pilot (#3).",
]:
    body(ws.cell(row=r,column=1), s); r+=1
r+=1

h1(ws.cell(row=r,column=1), "Integrated action plan", size=12); r+=1

def write_table(ws, r, headers, rows, widths):
    for j,htext in enumerate(headers, start=1):
        c = ws.cell(row=r, column=j, value=htext); c.font=HEADER_FONT; c.fill=HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    start = r
    r+=1
    for row in rows:
        for j,val in enumerate(row, start=1):
            c = ws.cell(row=r, column=j, value=val); c.font=BODY_FONT
            c.alignment = Alignment(wrap_text=True, vertical='top')
        r+=1
    for j,w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    return r

headers = ['#','Objective','Owner','KPI','Horizon','Risk','Source']
rows = [
 [1,'Ancillary fee attach push -- Self/app cart stage first (has existing attach to build on); Phlebo needs separate fix','Pratyush (CRM) / Tomar (product)','Ancillary rev/booking Rs 64 -> Rs 85+','3 weeks','Low -- no margin trade-off','Deepdive #1, sharpened by V4 Finding A'],
 [2,'Family add-on forcing -- discount structured on 2nd+ member, not flat on cart','Tomar (product)','Members/booking 1.22 -> 1.35+','30 days','Low','Deepdive #2, confirmed by V4 members/booking check'],
 [3,'Soft basket tiering pilot -- graduated VIP%, not a hard cliff, 2-3 cities only','Growth / Tomar','Blended ASP/booking','2-week pilot','Medium -- VIP-only segment already runs 70-77% GM%, protect it','Deepdive #3, adjusted from Gemini original'],
 [4,'Stacked-segment CPT investigation -- why CPT% rose in VIP+coupon bookings (24.9-26.7% range)','Mukesh (data)','Identify test/package driver','1 week','Low -- diagnostic only','Deepdive #4'],
 [5,'Fresh-member funnel check -- fresh share fell 30.8%->25.7% over 11 weeks','Kamal/Nikhil (acquisition)','Fresh% stabilize >=28%','4 weeks','Ties to broader retention issue','Deepdive #5'],
 [6,'HOLD 5 VIP-Gold-heavy price cuts until #4 resolves -- PL124I (48% VG share), PL08I (35%), PL334 (33%), PL91 (37%), PL203 (34%)','Pricing / Growth','No premature GM erosion on largest segment','Pending #4','Medium -- these are on this week\'s Price Action List','NEW -- V4 Finding B'],
]
r = write_table(ws, r, headers, rows, [4,34,20,22,12,26,26])
r+=1
h1(ws.cell(row=r,column=1), "Sequencing", size=12); r+=1
body(ws.cell(row=r,column=1), "Run #1 and #4 first -- pure analysis/config, zero customer-facing risk, and #1 alone plausibly recovers more ASP than any discount edit given the size of the ancillary drop. Apply #6 immediately (it's a hold, not new work). Only pilot #3 once #4 shows whether the stacked-segment margin issue is cost-side (leave discount alone) or genuinely discount-side (then tiering earns its place).")
r+=2
h1(ws.cell(row=r,column=1), "Open item from the deepdive worth chasing", size=12); r+=1
body(ws.cell(row=r,column=1), "ASP decline in the deepdive's WoW table starts in early-mid June, before the paid-VIP/MOV-removal transition (~1.5 months prior to mid-Aug). Jun1->Jun8 alone drops ASP/member -12%. The membership-switch story explains the July-August leg but not the full decline -- worth a quick look at what changed operationally in early June before attributing the whole slide to the paywall change.")

wb.save("/sessions/gallant-sharp-hopper/mnt/GM Improvement/GM_V4_Model_Aug2026.xlsx")
print("Saved. Sheets:", wb.sheetnames)
